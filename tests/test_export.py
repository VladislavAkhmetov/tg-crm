import unittest
import io
import openpyxl
import exporter

class ExportTest(unittest.TestCase):
    def setUp(self):
        self.mock_tasks = [
            {
                "id": 1,
                "sku": "SKU-001",
                "pool_name": "Кроссовки",
                "status": "completed",
                "assigned_to_id": 10001,
                "assigned_to_name": "Анна",
                "assigned_to_username": "anna_design",
                "assigned_at": "2026-10-01T10:00:00Z",
                "updated_at": "2026-10-02T15:30:00Z",
                "images_count": 6,
                "source_url": "https://cloud.example.com/sku-001",
                "internal_figma_url": "https://figma.com/file/1",
                "client_figma_url": "https://figma.com/file/2",
                "yandex_disk_url": "https://disk.yandex.ru/d/123",
                "bitrix_url": "https://b24.example.com/tasks/123",
                "notes": "Правка по цвету",
                "is_archived": 0
            },
            {
                "id": 2,
                "sku": "SKU-002",
                "pool_name": "Кроссовки",
                "status": "in_progress",
                "assigned_to_id": 10002,
                "assigned_to_name": "Михаил",
                "assigned_to_username": "misha_art",
                "assigned_at": "2026-10-03T09:00:00Z",
                "updated_at": "2026-10-03T11:00:00Z",
                "images_count": 8,
                "notes": "",
                "is_archived": 0
            },
            {
                "id": 3,
                "sku": "SKU-003",
                "pool_name": "Куртки",
                "status": "available",
                "assigned_to_id": None,
                "assigned_to_name": None,
                "assigned_to_username": None,
                "assigned_at": None,
                "updated_at": "2026-10-01T12:00:00Z",
                "images_count": None,
                "notes": "",
                "is_archived": 0
            }
        ]

    def test_excel_generation_single_pool(self):
        single_pool_tasks = [t for t in self.mock_tasks if t["pool_name"] == "Кроссовки"]
        excel_bytes = exporter.generate_excel(single_pool_tasks, pool_names=["Кроссовки"])
        self.assertIsInstance(excel_bytes, bytes)
        self.assertGreater(len(excel_bytes), 1000)

        wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
        sheet_names = wb.sheetnames
        self.assertIn("Сводка", sheet_names)
        self.assertIn("Кроссовки", sheet_names)

        # Check summary sheet content
        ws_sum = wb["Сводка"]
        self.assertEqual(ws_sum["A1"].value, "ОТЧЕТ ПО ЗАДАЧАМ И ПУЛАМ")

        # Check pool sheet content
        ws_pool = wb["Кроссовки"]
        headers = [cell.value for cell in ws_pool[1]]
        self.assertIn("Артикул (SKU)", headers)
        self.assertIn("Дизайнер", headers)
        self.assertIn("Изображений (шт.)", headers)
        self.assertIn("Задача в Битрикс24", headers)

        # Check row 2 (first task)
        row2_values = [cell.value for cell in ws_pool[2]]
        self.assertIn("SKU-001", row2_values)
        self.assertIn("Готово", row2_values)
        self.assertIn("Анна (@anna_design)", row2_values)
        self.assertIn(6, row2_values)

    def test_excel_generation_multiple_pools(self):
        excel_bytes = exporter.generate_excel(self.mock_tasks)
        wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
        sheet_names = wb.sheetnames
        self.assertIn("Сводка", sheet_names)
        self.assertIn("Все задачи", sheet_names)
        self.assertIn("Кроссовки", sheet_names)
        self.assertIn("Куртки", sheet_names)

    def test_excel_generation_empty_tasks(self):
        excel_bytes = exporter.generate_excel([])
        self.assertIsInstance(excel_bytes, bytes)
        wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
        self.assertIn("Сводка", wb.sheetnames)

    def test_google_sheets_payload_structure(self):
        payload = exporter.prepare_google_sheets_payload(self.mock_tasks)
        self.assertIn("generated_at", payload)
        self.assertIn("sheets", payload)
        self.assertGreater(len(payload["sheets"]), 1)

        summary_sheet = next(s for s in payload["sheets"] if s["title"] == "Сводка")
        self.assertTrue(any("Всего задач (артикулов)" in str(row) for row in summary_sheet["rows"]))

        pool_sheet = next(s for s in payload["sheets"] if s["title"] == "Кроссовки")
        self.assertIn("SKU-001", [row[1] for row in pool_sheet["rows"][1:]])

    def test_lead_time_calculation(self):
        lead = exporter.calculate_lead_time("2026-10-01T10:00:00Z", "2026-10-02T15:30:00Z")
        self.assertEqual(lead, "1 д. 5 ч.")
        self.assertEqual(exporter.calculate_lead_time(None, None), "—")

    def test_sanitize_sheet_title(self):
        self.assertEqual(exporter.sanitize_sheet_title("Пул: [1] / 2 * 3?"), "Пул_ _1_ _ 2 _ 3_")
        self.assertLessEqual(len(exporter.sanitize_sheet_title("Очень длинное название пула с деталями и подробностями")), 28)

if __name__ == "__main__":
    unittest.main()
